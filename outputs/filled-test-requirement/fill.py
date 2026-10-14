from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import hashlib

source=Path(r'd:\Downloads\TEST REQUIREMENT DOCUMENT.docx')
out=Path(__file__).parent
d=Document(source)
answers={
0:{1:'Zea Voice — AI voice-agent platform. Initial testing configuration: Balaji Robotic Rehab Hospital.',2:'Phase 1 / Beta 1. Deployed build and Git commit: to be confirmed.',3:'Shanmugam / Zea Voice development team; full contributor list to be confirmed.',4:'To be confirmed by Project Manager.',5:'To be confirmed.',6:'HIGH. Critical dialing, permission, billing and false action-confirmation defects block release.'},
1:{1:'Automate incoming/outgoing enquiries, approved knowledge answers, follow-ups and company-scoped conversation history.',2:'Reduces missed enquiries, repetitive calling work, inconsistent follow-up and scattered call records.',3:'Company Developers, Company Users, Super Admin and customers/patients calling the configured agents.',4:'Super Admin, Company Developer and Company User. Enabled tabs remain subject to role permissions and company isolation.'},
2:{1:'Login/roles/tab access; agent configuration saves; inbound/outbound audio and interruption; KB/live prices; contact/history; appointment-request tools; callbacks; CSV import; schedules/retries; company queues; reports, recording and billing; share links.',2:'WhatsApp calling; unattached staff/CRM/transfer tools; additional provider/language combinations; concurrency above approved Beta capacity. Confirm scope with PM.',3:'Client exclusions are not established. PM/client must identify and approve any additional features.',5:'Post-tool context, contact reuse, callbacks, TTS recovery, report permissions, account tab access and permanent campaign deletion. Verify deployment before retest.',6:'Conversation layout and summaries; tab-access toggles; typed DELETE warning with task counts; User cost/source visibility.',7:'Retest reported fixes: lost appointment context after contact update; duplicate speech; relative callback rejection; punctuation-only TTS requests; deleted-campaign queue leftovers. Live verification required.',8:'Reuse known details; ask missing fields only; verify tool success; schedule callbacks through queues; apply saved UI controls; guard permanent campaign/task deletion.',9:'Review report contact-name mapping, retrieval quality, provider diagnostics, alerts, restart recovery and larger load tests. Confirm backlog with PM.'},
4:{1:'YES. Record the theme for each defect and test implemented light/dark modes. Colourful mode is not confirmed; mark Not Applicable unless provided. Avoid duplicate issues for the same root cause.',2:'Evidence required: test ID, steps, expected/actual result, role/company, theme, IST time, Call/Campaign/Task ID, screenshots and sanitized logs.',3:'Release gate: resolve critical defects; verify migrations, backup/rollback, provider limits and task/call recovery.',4:'Use approved test numbers and synthetic patient data. Do not include passwords, API keys or real medical records in bug reports.'},
5:{1:'Website: https://voice.zeacrm.com/\nDeveloper: support@balajiroboticrehab.com\nCompany User: dr.nithya@balajiroboticrehab.com\nConfirm active access; share current passwords separately through an approved secure channel.',2:'Appointment intake interrupted by a fee question; returning contact; callback after 2 minutes; reschedule/cancel; 100 unique authorized leads with company/campaign concurrency 2, inbound queue 100 and outbound queue 1000.',3:'CSV headers: name,phone,remarks; full E.164 numbers as text; team-owned consented contacts; verified KB facts/current fees; explicit caller name corrections. Repeated phone numbers are duplicates, not separate leads.',4:'Blank/malformed or scientific-notation phone; duplicate phone; missing CSV header; ambiguous/past time; unavailable tool; expired/revoked link; disabled tab; cross-company ID; insufficient credits.'}}
def write(cell,text):
 p=cell.paragraphs[0]
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 for extra in list(cell.paragraphs)[1:]:extra._p.getparent().remove(extra._p)
 run=p.add_run(text);run.font.size=Pt(10)
 p.paragraph_format.space_after=Pt(4)
 p.paragraph_format.space_before=Pt(4)
for ti,rows in answers.items():
 for ri,text in rows.items():
  if ti==4 and ri>1:write(d.tables[ti].rows[ri].cells[0],['','', 'Failure evidence','Release checks','Privacy and test safety'][ri])
  write(d.tables[ti].rows[ri].cells[1],text)
matrix=[
['Features','Super Admin','Company Developer','Company User'],
['Manage companies/providers/credits','YES','NO','NO'],
['Manage accounts and tab access','YES','NO','NO'],
['Dashboard/Analytics/Reports/Conversations','YES within authorized scope','YES if tab enabled','YES if tab enabled'],
['Edit agents and manage campaigns/queue','YES within company context','YES if tab enabled','NO; read-only where exposed'],
['Report Call Cost and Answer Sources','YES','YES with report access','NO'],
['API Keys/Settings/Integrations','YES where applicable','YES if tab enabled','NO']]
for row,values in zip(d.tables[3].rows,matrix):
 for cell,text in zip(row.cells,values):write(cell,text)
for t in d.tables:
 for r in t.rows:
  pr=r._tr.get_or_add_trPr()
  for h in list(pr.findall(qn('w:trHeight'))):pr.remove(h)
  no=OxmlElement('w:cantSplit');pr.append(no)
 # Repeat the existing header when a table spans pages.
 pr=t.rows[0]._tr.get_or_add_trPr();hdr=OxmlElement('w:tblHeader');pr.append(hdr)
out.mkdir(parents=True,exist_ok=True)
(out/'artifact.md').write_text('Reference: '+str(source)+'\nSHA256: '+hashlib.sha256(source.read_bytes()).hexdigest()+'\nPreserve six section headings, six tables, page geometry and existing style. Fill answer cells and role matrix. Replace blank additional-requirement labels with meaningful fields. Allow rows to expand; repeat table headers. Unknown owners, dates and build remain unconfirmed.\n',encoding='utf-8')
d.save(out/'Zea_Voice_Test_Requirement_Filled.docx')
print('Saved filled document; six sections retained.')
